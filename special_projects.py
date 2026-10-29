"""Validated, manually imported 4G transport snapshots, independent of app state.

Call init_store(connect, sql) during application initialization. Imports return the
same snapshot shape as read_snapshot and raise ValueError on invalid workbooks.
Formula cells require saved Excel cached results; this module never evaluates them.
"""

import hashlib
import io
import json
import unicodedata
from datetime import date, datetime, timezone
from contextlib import contextmanager
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook


WORKBOOK_URL = (
    'https://vobissgh1-my.sharepoint.com/:x:/r/personal/andrew_annor_vobissgh_com/'
    '_layouts/15/Doc.aspx?sourcedoc=%7BFC7AE9C4-0679-4CE1-B5B3-3F31A0FCC5D2%7D'
    '&file=4G%20Transport%20link%20plan_19.10.26.xlsx'
    '&fromShare=true&action=default&mobileredirect=true'
)
MAX_COMPRESSED = 20 * 1024 * 1024
MAX_UNCOMPRESSED = 40 * 1024 * 1024
MAX_LINKS = 10000
MAX_SCANNED_ROWS = 20000
MAX_SCANNED_COLUMNS = 256
STAGE_NAMES = (
    'Cable Stringing', 'Cable Entry - Site A', 'Cable Entry - Site B',
    'Fiber Splicing', 'ODF Termination - Site A', 'ODF Termination - Site B',
    'OTDR / Fiber Test', 'Nokia Install - Site A', 'Nokia Install - Site B',
    'Power-up & Commissioning', 'Provisioning', 'Service Up (E2E Test)',
)
STATUSES = ('Not Started', 'In Progress', 'Completed', 'Blocked', 'N/A')
FIELDS = {
    'link': ('Link', 500), 'a': ('Site A', 250), 'b': ('Site B', 250),
    'type': ('Link Type', 100), 'priority': ('Priority', 100),
    'team': ('Responsible', 250), 'fiber': ('Fiber Link', 500),
    'issue': ('Current Blocker / Issue', 4000), 'remarks': ('Remarks', 4000),
    'updated': ('Last Updated', 100),
}


def _normalize(value):
    text = unicodedata.normalize('NFKC', str(value or '')).casefold()
    return ''.join(char for char in text if char.isalnum())


def _default():
    return dict(title='4G Transport', workbook_url=WORKBOOK_URL, imported_at=None,
                connection='not_configured', links=[])


@contextmanager
def _connection(connect):
    conn = connect()
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def init_store(connect, sql):
    """Create the isolated snapshot table; does not seed a source row."""
    with _connection(connect) as conn:
        conn.execute(sql('''CREATE TABLE IF NOT EXISTS special_project_sources (
            source_key TEXT PRIMARY KEY,
            snapshot_json TEXT NOT NULL,
            imported_at TEXT NOT NULL
        )'''))


def read_snapshot(connect, sql):
    """Read the persisted 4g source, or return a fresh unconfigured snapshot."""
    with _connection(connect) as conn:
        row = conn.execute(sql('''SELECT snapshot_json FROM special_project_sources
                                  WHERE source_key = ?'''), ('4g',)).fetchone()
    if row is None:
        return _default()
    value = row['snapshot_json'] if hasattr(row, 'keys') else row[0]
    return json.loads(value)


def _checked_archive(payload):
    if not isinstance(payload, bytes) or not payload:
        raise ValueError('A nonempty XLSX byte payload is required')
    if len(payload) > MAX_COMPRESSED:
        raise ValueError('Workbook exceeds the 20 MB compressed limit')
    try:
        with ZipFile(io.BytesIO(payload)) as archive:
            members = archive.infolist()
            if sum(item.file_size for item in members) > MAX_UNCOMPRESSED:
                raise ValueError('Workbook exceeds the 40 MB uncompressed limit')
            if any(item.flag_bits & 1 for item in members):
                raise ValueError('Encrypted workbooks are not supported')
            # Check actual inflated sizes as well as the untrusted ZIP directory.
            total = 0
            for item in members:
                with archive.open(item) as stream:
                    while True:
                        chunk = stream.read(64 * 1024)
                        if not chunk:
                            break
                        total += len(chunk)
                        if total > MAX_UNCOMPRESSED:
                            raise ValueError('Workbook exceeds the uncompressed limit')
                if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
                    _check_coordinates(archive, item)
    except (BadZipFile, RuntimeError, NotImplementedError, OSError) as exc:
        raise ValueError('Invalid XLSX archive') from exc


def _check_coordinates(archive, item):
    """Reject sparse coordinates before openpyxl synthesizes empty grid cells."""
    try:
        with archive.open(item) as stream:
            for _, element in ElementTree.iterparse(stream, events=('end',)):
                tag = element.tag.rsplit('}', 1)[-1]
                coordinate = element.get('r', '')
                if tag == 'row' and coordinate and int(coordinate) > MAX_SCANNED_ROWS:
                    raise ValueError('Workbook exceeds safe worksheet limits')
                if tag == 'c' and coordinate:
                    letters = ''.join(char for char in coordinate if char.isalpha()).upper()
                    digits = ''.join(char for char in coordinate if char.isdigit())
                    column = 0
                    for char in letters:
                        column = column * 26 + ord(char) - ord('A') + 1
                    if column > MAX_SCANNED_COLUMNS or (digits and int(digits) > MAX_SCANNED_ROWS):
                        raise ValueError('Workbook exceeds safe worksheet limits')
                element.clear()
    except ElementTree.ParseError as exc:
        raise ValueError('Invalid worksheet XML') from exc


def _text(value, limit, row_number, field):
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    text = str(value).strip() if value is not None else ''
    if len(text) > limit:
        raise ValueError(f'Row {row_number}: {field} exceeds {limit} characters')
    return text


def _header(rows):
    required = {_normalize(name) for name in ('Link', *STAGE_NAMES)}
    known = required | {_normalize(name) for name, _ in FIELDS.values()} | {'no'}
    for number, row in enumerate(rows, 1):
        if number > 10:
            break
        normalized = tuple(_normalize(value) for value in row)
        if required.issubset(normalized):
            if any(normalized.count(name) > 1 for name in known):
                raise ValueError(f'Row {number}: duplicate header columns')
            return number, {name: index for index, name in enumerate(normalized) if name}
    raise ValueError('First 10 rows must contain Link and all 12 stage headers')


def _parse_row(row, columns, number):
    def cell(name):
        index = columns.get(_normalize(name))
        return row[index] if index is not None and index < len(row) else None

    if cell('Link') is None or not str(cell('Link')).strip():
        return None
    fields = {key: _text(cell(name), limit, number, name)
              for key, (name, limit) in FIELDS.items()}
    stage_lookup = {_normalize(status): status for status in STATUSES}
    stages = []
    for name in STAGE_NAMES:
        raw = _text(cell(name), 100, number, name)
        status = stage_lookup.get(_normalize(raw))
        if status is None:
            raise ValueError(f'Row {number}: missing or invalid status for {name}')
        stages.append(status)
    completed = stages.count('Completed')
    applicable = len(stages) - stages.count('N/A')
    if 'Blocked' in stages:
        overall = 'Blocked'
    elif applicable and completed == applicable:
        overall = 'Completed'
    elif completed or 'In Progress' in stages:
        overall = 'In Progress'
    else:
        overall = 'Not Started'
    identity_no = _text(cell('No'), 100, number, 'No')
    identity = ('no' if identity_no else 'link',
                ' '.join((identity_no or fields['link']).split()).casefold(),
                ' '.join(fields['type'].split()).casefold())
    identifier = hashlib.sha256(json.dumps(identity).encode('utf-8')).hexdigest()
    return dict(id=identifier, **fields, stages=stages, status=overall,
                fraction=completed / applicable if applicable else 0)


def _parse_links(payload):
    _checked_archive(payload)
    try:
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True,
                           keep_links=False)
    except Exception as exc:
        raise ValueError('Unable to read XLSX workbook') from exc
    try:
        if 'Link Status Tracker' not in wb.sheetnames:
            raise ValueError('Missing Link Status Tracker worksheet')
        ws = wb['Link Status Tracker']
        # Ignore potentially stale worksheet dimensions recorded by Excel writers.
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        header_number, columns = _header(rows)
        links = []
        identities = set()
        for number, row in enumerate(rows, header_number + 1):
            if number > MAX_SCANNED_ROWS or len(row) > MAX_SCANNED_COLUMNS:
                raise ValueError('Workbook exceeds safe worksheet limits')
            item = _parse_row(row, columns, number)
            if item is None:
                continue
            if item['id'] in identities:
                raise ValueError(f'Row {number}: duplicate link identity')
            if len(links) >= MAX_LINKS:
                raise ValueError(f'Workbook exceeds {MAX_LINKS} links')
            identities.add(item['id'])
            links.append(item)
        if not links:
            raise ValueError('Workbook contains no valid links')
        return links
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Invalid worksheet contents') from exc
    finally:
        wb.close()


def import_snapshot(payload, connect, sql):
    """Validate XLSX bytes and atomically replace source 4g; no partial imports.

    stages is a list of 12 canonical status strings in STAGE_NAMES order.
    connection='imported' denotes a manual snapshot, not a live connection.
    Database exceptions propagate; the connection context rolls back failed writes.
    """
    links = _parse_links(payload)
    timestamp = datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z')
    snapshot = dict(_default(), imported_at=timestamp, connection='imported', links=links)
    encoded = json.dumps(snapshot, ensure_ascii=True, allow_nan=False)
    with _connection(connect) as conn:
        conn.execute(sql('''INSERT INTO special_project_sources
            (source_key, snapshot_json, imported_at) VALUES (?, ?, ?)
            ON CONFLICT (source_key) DO UPDATE SET
                snapshot_json = excluded.snapshot_json,
                imported_at = excluded.imported_at'''), ('4g', encoded, timestamp))
    return snapshot
